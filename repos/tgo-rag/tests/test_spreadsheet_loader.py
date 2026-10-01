import tempfile
import unittest
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook

from src.rag_service.spreadsheet_loader import extract_spreadsheet


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class SpreadsheetLoaderTest(unittest.TestCase):
    def test_xlsx_preserves_all_sheets_coordinates_values_and_formulas(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "知识库.xlsx"
            book = Workbook()
            products = book.active
            products.title = "产品"
            products.append(["名称", "价格"])
            products.append(["K9", 100])
            products.append(["K8", 80])
            products["B4"] = "=SUM(B2:B3)"
            products.merge_cells("A6:B6")
            products["A6"] = "备注"
            hidden = book.create_sheet("内部参数")
            hidden.sheet_state = "hidden"
            hidden["C2"] = "不可丢失"
            book.save(path)

            sheets = extract_spreadsheet(str(path), XLSX_MIME)

        self.assertEqual([sheet.name for sheet in sheets], ["产品", "内部参数"])
        self.assertIn("A1: 名称 | B1: 价格", sheets[0].content)
        self.assertIn("B4: =SUM(B2:B3)", sheets[0].content)
        self.assertIn("合并单元格: A6:B6", sheets[0].content)
        self.assertIn("C2: 不可丢失", sheets[1].content)
        self.assertEqual(sheets[1].metadata["sheet_state"], "hidden")


if __name__ == "__main__":
    unittest.main()
