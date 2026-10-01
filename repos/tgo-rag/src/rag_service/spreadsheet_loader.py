"""Loss-minimizing Excel extraction for knowledge-base ingestion."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS_MIME = "application/vnd.ms-excel"


@dataclass(frozen=True)
class SpreadsheetSheet:
    name: str
    content: str
    metadata: dict[str, Any]


def _format_value(value: Any) -> str:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _column_label(index: int) -> str:
    label = ""
    current = index + 1
    while current:
        current, remainder = divmod(current - 1, 26)
        label = chr(65 + remainder) + label
    return label


def _extract_xlsx(file_path: str) -> list[SpreadsheetSheet]:
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=False, data_only=False)
    sheets: list[SpreadsheetSheet] = []
    try:
        for worksheet in workbook.worksheets:
            lines = [f"工作表: {worksheet.title}"]
            merged_ranges = [str(value) for value in worksheet.merged_cells.ranges]
            if merged_ranges:
                lines.append("合并单元格: " + ", ".join(merged_ranges))
            for row in worksheet.iter_rows():
                cells = [
                    f"{cell.coordinate}: {_format_value(cell.value)}"
                    for cell in row
                    if cell.value is not None
                ]
                if cells:
                    lines.append(" | ".join(cells))
            sheets.append(SpreadsheetSheet(
                name=worksheet.title,
                content="\n".join(lines),
                metadata={
                    "sheet_name": worksheet.title,
                    "sheet_state": worksheet.sheet_state,
                    "source": str(Path(file_path)),
                },
            ))
    finally:
        workbook.close()
    return sheets


def _extract_xls(file_path: str) -> list[SpreadsheetSheet]:
    import xlrd

    workbook = xlrd.open_workbook(file_path, formatting_info=False)
    sheets: list[SpreadsheetSheet] = []
    for worksheet in workbook.sheets():
        lines = [f"工作表: {worksheet.name}"]
        for row_index in range(worksheet.nrows):
            cells = []
            for column_index in range(worksheet.ncols):
                value = worksheet.cell_value(row_index, column_index)
                if value not in (None, ""):
                    coordinate = f"{_column_label(column_index)}{row_index + 1}"
                    cells.append(f"{coordinate}: {_format_value(value)}")
            if cells:
                lines.append(" | ".join(cells))
        sheets.append(SpreadsheetSheet(
            name=worksheet.name,
            content="\n".join(lines),
            metadata={
                "sheet_name": worksheet.name,
                "sheet_state": "visible",
                "source": str(Path(file_path)),
            },
        ))
    return sheets


def extract_spreadsheet(file_path: str, content_type: str) -> list[SpreadsheetSheet]:
    """Extract every worksheet without row, sheet, or cell-count truncation."""
    suffix = Path(file_path).suffix.lower()
    if content_type == XLSX_MIME or suffix == ".xlsx":
        return _extract_xlsx(file_path)
    if content_type == XLS_MIME or suffix == ".xls":
        return _extract_xls(file_path)
    raise ValueError(f"Unsupported spreadsheet type: {content_type or suffix}")
